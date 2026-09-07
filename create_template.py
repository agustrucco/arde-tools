"""
Generate template.xlsx with sample fictional data.
Run once to create the workbook: python create_template.py
"""

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

DARK_FILL  = PatternFill("solid", fgColor="2D2D2D")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GRAY_FILL  = PatternFill("solid", fgColor="F5F5F5")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
NORMAL_FONT = Font(name="Calibri")

def header(ws, row, cols):
    for i, col in enumerate(cols, 1):
        c = ws.cell(row, i, col)
        c.fill = DARK_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")

wb = Workbook()

# ── STOCK ──────────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = "STOCK"

ws.cell(1, 1, "#")
header(ws, 1, ["#", "PIEZA", "MODELO", "METAL",
               "STOCK INICIAL", "VENDIDOS", "PERDIDOS", "TALLER",
               "STOCK LOCAL", "ESTADO", "PRECIO BASE"])

sample_stock = [
    (1, "ANILLO", "LUNA", "ALPACA", 20, 0, 0, 0),
    (2, "ANILLO", "LUNA", "BRONCE", 15, 0, 0, 0),
    (3, "ARO",    "SOL",  "ALPACA", 30, 0, 0, 0),
    (4, "ARO",    "SOL",  "BRONCE", 25, 0, 0, 0),
    (5, "DIJE",   "ESTRELLA", "ALPACA", 40, 0, 0, 0),
    (6, "COLLAR", "CADENA FINA", "ALPACA", 10, 0, 1, 0),
]

for i, (num, pieza, modelo, metal, ini, vend, perd, tall) in enumerate(sample_stock, 2):
    ws.cell(i, 1, num)
    ws.cell(i, 2, pieza)
    ws.cell(i, 3, modelo)
    ws.cell(i, 4, metal)
    ws.cell(i, 5, ini)
    ws.cell(i, 6, f"=COUNTIFS('ENERO 26'!D:D,B{i},'ENERO 26'!E:E,C{i},'ENERO 26'!F:F,D{i},'ENERO 26'!J:J,\"SI\")")
    ws.cell(i, 7, perd)
    ws.cell(i, 8, tall)
    ws.cell(i, 9, f"=IF(E{i}=\"\",\"\",E{i}-F{i}-IF(G{i}=\"\",0,G{i})-IF(H{i}=\"\",0,H{i}))")
    stock_local = ini - vend - perd - tall
    if stock_local <= 0:
        estado = "SIN STOCK"
    elif stock_local <= 3:
        estado = "STOCK BAJO"
    else:
        estado = "OK"
    ws.cell(i, 10, estado).fill = GREEN_FILL
    ws.cell(i, 11, 2500)

ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 16
ws.column_dimensions["J"].width = 18

# ── ENERO 26 (sample sales sheet) ─────────────────────────────────────────────
ws2 = wb.create_sheet("ENERO 26")
header(ws2, 1, ["F. PEDIDO", "ORDEN", "PDV", "PIEZA", "MODELO", "METAL",
                "F. PAGO", "PRECIO", "FORMA PAGO", "DESC STOCK",
                "F. ENTREGA", "ESTADO", "CLIENTE", "CELU", "VENDEDOR/A"])

sample_sales = [
    ("2026-01-05", "#001", "LOCAL",     "ANILLO", "LUNA",     "ALPACA", "2026-01-05", 2500, "GETNET",       "SI", "2026-01-05", "ARCHIVADO", "Maria Garcia",   "1155550001", "VENDEDOR A"),
    ("2026-01-08", "#002", "INSTAGRAM", "ARO",    "SOL",      "BRONCE", "2026-01-08", 2000, "MERCADO PAGO", "SI", "2026-01-09", "ARCHIVADO", "Laura Perez",    "1155550002", "VENDEDOR B"),
    ("2026-01-12", "#003", "ECOMMERCE", "DIJE",   "ESTRELLA", "ALPACA", "2026-01-12", 1800, "TRANSFERENCIA","SI", "2026-01-14", "ARCHIVADO", "Ana Rodriguez",  "1155550003", "VENDEDOR A"),
    ("2026-01-20", "-",    "LOCAL",     "ANILLO", "LUNA",     "BRONCE", "2026-01-20", 2000, "EFECTIVO",     "SI", "2026-01-20", "ARCHIVADO", "Camila Lopez",   "1155550004", "VENDEDOR B"),
]

for i, row_data in enumerate(sample_sales, 2):
    for j, val in enumerate(row_data, 1):
        ws2.cell(i, j, val)

for col in ["A", "G", "K"]:
    ws2.column_dimensions[col].width = 13
ws2.column_dimensions["M"].width = 16

# ── TAREAS ─────────────────────────────────────────────────────────────────────
ws3 = wb.create_sheet("TAREAS")
header(ws3, 1, ["#", "TAREA", "DETALLE", "RESPONSABLE", "ESTADO", "PRIORIDAD", "", "RESPUESTA"])
ws3.cell(2, 1, 1)
ws3.cell(2, 2, "Update STOCK INICIAL")
ws3.cell(2, 3, "Do a physical count and adjust column E in STOCK")
ws3.cell(2, 4, "Staff")
ws3.cell(2, 5, "PENDIENTE")
ws3.cell(2, 6, "ALTA")

ws3.cell(3, 1, 2)
ws3.cell(3, 2, "Add new month")
ws3.cell(3, 3, "Run: python inventory_tools.py add_month 'FEB 26'")
ws3.cell(3, 4, "Operator")
ws3.cell(3, 5, "LISTO")
ws3.cell(3, 6, "MEDIA")
ws3.cell(3, 8, "Sheet created and ready.")

# ── ENCARGOS ───────────────────────────────────────────────────────────────────
ws4 = wb.create_sheet("ENCARGOS")
header(ws4, 1, ["DESCRIPCION", "CERA", "CLIENTE", "CELU", "PRECIO", "ESTADO"])
ws4.cell(2, 1, "Ring size 18 LUNA ALPACA model")
ws4.cell(2, 2, "POR HACER")
ws4.cell(2, 3, "Valentina Sosa")
ws4.cell(2, 4, "1155550099")
ws4.cell(2, 5, 3000)
ws4.cell(2, 6, "PENDIENTE")

wb.save("template.xlsx")
print("Template created: template.xlsx")