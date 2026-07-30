#!/usr/bin/env python3
"""
arde_tools.py — Utilidades para NUEVO ARDE AGUSTIN.xlsx

Uso desde terminal:
    python arde_tools.py styling              # aplicar colores a todas las hojas
    python arde_tools.py fix_formulas         # regenerar fórmulas VENDIDOS en STOCK
    python arde_tools.py add_month "SEP 26"   # crear nueva hoja de mes
    python arde_tools.py stock_report         # ver estado del stock por consola
    python arde_tools.py fix_desc_stock       # convertir TRUE/FALSE a SI/NO
    python arde_tools.py validate             # reportar entradas con posibles typos
    python arde_tools.py validate --fix       # reportar Y corregir automáticamente
"""

import sys
import difflib
from pathlib import Path
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import column_index_from_string

FILE_PATH = Path(__file__).parent / "NUEVO ARDE AGUSTIN.xlsx"

# ── Colores ────────────────────────────────────────────────────────────────────

DARK_FILL   = PatternFill("solid", fgColor="2D2D2D")
GRAY_FILL   = PatternFill("solid", fgColor="F5F5F5")
WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")
GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
RED_FILL    = PatternFill("solid", fgColor="FFC7CE")
LGRAY_FILL  = PatternFill("solid", fgColor="E0E0E0")

HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
NORMAL_FONT = Font(name="Calibri")

# ── Mapeo de columnas ──────────────────────────────────────────────────────────

# STOCK (estructura post-julio 2026)
STOCK_COLS = {
    "num":          "A",
    "pieza":        "B",
    "modelo":       "C",
    "metal":        "D",
    "stock_inicial":"E",
    "vendidos":     "F",
    "perdidos":     "G",
    "taller":       "H",
    "stock_local":  "I",
    "estado":       "J",
    "precio_base":  "K",
}

# Columna DESC STOCK por hoja (verificado contra el archivo real)
DESC_STOCK_COL = {
    "ENERO26":   "I",
    "FEBRERO26": "I",
    "MARZO26":   "I",
    "ABRIL26":   "J",
    "MAYO26":    "J",
    "JUNIO 26":  "J",
    "JULIO 26":  "J",
    "AGOSTO 26": "J",
}

MONTHLY_SHEETS = [
    "ENERO26", "FEBRERO26", "MARZO26", "ABRIL26",
    "MAYO26", "JUNIO 26", "JULIO 26", "AGOSTO 26",
]

# Hojas que NO son de ventas (excluir de la deteccion de meses)
NON_MONTH_SHEETS = {
    "STOCK", "README", "TAREAS", "PROPUESTAS", "DASHBOARD",
    "ENCARGOS", "ZIEZTA", "PAUYSOL ABRIL",
}

ESTADO_STOCK_COLORS = {
    "OK":                   GREEN_FILL,
    "STOCK BAJO":           YELLOW_FILL,
    "SIN STOCK":            RED_FILL,
    "Cargar stock inicial": LGRAY_FILL,
}

ESTADO_TAREAS_COLORS = {
    "PENDIENTE":   RED_FILL,
    "EN PROGRESO": YELLOW_FILL,
    "LISTO":       GREEN_FILL,
}

# ── Helpers internos ───────────────────────────────────────────────────────────

def _load(data_only=False):
    return load_workbook(FILE_PATH, data_only=data_only)

def _save(wb):
    wb.save(FILE_PATH)
    print(f"  Guardado: {FILE_PATH.name}")

def _find_header_row(ws, marker="#"):
    """Devuelve el número de fila donde la col A tiene el valor marker."""
    for r in range(1, 6):
        if ws.cell(r, 1).value == marker:
            return r
    return 1

def _col(name):
    return column_index_from_string(STOCK_COLS[name])

def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = DARK_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

def _style_data_row(ws, row, ncols, even):
    bg = GRAY_FILL if even else WHITE_FILL
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
        if rgb in (None, "00000000", "FFFFFFFF", "FFF5F5F5", "FFF0F0F0"):
            cell.fill = bg
        cell.font = NORMAL_FONT

def _add_desc_stock_dv(ws, col_letter, first_row=2, last_row=1000):
    dv = DataValidation(
        type="list",
        formula1='"SI,NO"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.sqref = f"{col_letter}{first_row}:{col_letter}{last_row}"
    ws.add_data_validation(dv)

def _month_sheets(wb):
    """Devuelve las hojas de ventas mensuales presentes, en orden del libro.

    Incluye cualquier hoja nueva tipo 'SEP 26' / 'AGOSTO 26' automaticamente:
    se considera hoja de mes si termina en un año de dos dígitos ('26', '27'...)
    y no está en NON_MONTH_SHEETS.
    """
    import re
    months = []
    for name in wb.sheetnames:
        if name in NON_MONTH_SHEETS:
            continue
        if name in DESC_STOCK_COL or re.search(r"\b2\d\b", name):
            months.append(name)
    # dedupe preservando orden
    seen = set()
    return [m for m in months if not (m in seen or seen.add(m))]

def _stock_piezas(wb):
    """Lista de tipos de PIEZA únicos en STOCK, en orden de aparición."""
    if "STOCK" not in wb.sheetnames:
        return []
    st = wb["STOCK"]
    hr = _find_header_row(st)
    seen, out = set(), []
    r = hr + 1
    while st.cell(r, _col("pieza")).value:
        v = st.cell(r, _col("pieza")).value
        if v and v not in seen:
            seen.add(v); out.append(v)
        r += 1
    return out

# ── Función: apply_styling ─────────────────────────────────────────────────────

def apply_styling(wb=None):
    """Aplica colores consistentes a STOCK, TAREAS y todas las hojas mensuales."""
    own = wb is None
    if own:
        wb = _load()

    # STOCK
    ws = wb["STOCK"]
    ncols = _col("precio_base")
    hr = _find_header_row(ws)
    _style_header(ws, hr, ncols)
    r, even = hr + 1, True
    while ws.cell(r, _col("pieza")).value:
        _style_data_row(ws, r, ncols, even)
        estado = ws.cell(r, _col("estado")).value
        if estado in ESTADO_STOCK_COLORS:
            ws.cell(r, _col("estado")).fill = ESTADO_STOCK_COLORS[estado]
        r += 1
        even = not even

    # TAREAS
    ws = wb["TAREAS"]
    ncols = 8
    hr = _find_header_row(ws)
    _style_header(ws, hr, ncols)
    r, even = hr + 1, True
    while ws.cell(r, 1).value:
        _style_data_row(ws, r, ncols, even)
        estado = ws.cell(r, 5).value  # col E = ESTADO
        if estado in ESTADO_TAREAS_COLORS:
            ws.cell(r, 5).fill = ESTADO_TAREAS_COLORS[estado]
        r += 1
        even = not even

    # Hojas mensuales
    all_months = MONTHLY_SHEETS + [s for s in wb.sheetnames
                                    if s not in MONTHLY_SHEETS and "26" in s
                                    and s not in ("PAUYSOL ABRIL",)]
    for name in all_months:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        ncols = ws.max_column
        if ncols < 4:
            continue
        _style_header(ws, 1, ncols)
        r, even = 2, True
        while ws.cell(r, 1).value:
            _style_data_row(ws, r, ncols, even)
            r += 1
            even = not even

    if own:
        _save(wb)
        print("  Styling aplicado.")

# ── Función: fix_vendidos_formulas ─────────────────────────────────────────────

def fix_vendidos_formulas(wb=None):
    """Regenera las fórmulas VENDIDOS en STOCK para TODAS las hojas de mes.

    Incluye automáticamente AGOSTO 26 y cualquier mes nuevo que exista en el
    archivo (detectado por _month_sheets). La columna DESC STOCK de cada mes se
    toma de DESC_STOCK_COL (por defecto 'J' para meses no listados).
    """
    own = wb is None
    if own:
        wb = _load()

    ws = wb["STOCK"]
    hr = _find_header_row(ws)

    months = _month_sheets(wb)

    def _quote(name):
        return f"'{name}'" if " " in name else name

    def formula(row):
        parts = []
        for name in months:
            ds_col = DESC_STOCK_COL.get(name, "J")
            q = _quote(name)
            parts.append(
                f'COUNTIFS({q}!D:D,B{row},{q}!E:E,C{row},'
                f'{q}!F:F,D{row},{q}!{ds_col}:{ds_col},"SI")'
            )
        return "=" + "+".join(parts)

    vendidos_col = _col("vendidos")
    r = hr + 1
    count = 0
    while ws.cell(r, _col("pieza")).value:
        ws.cell(r, vendidos_col).value = formula(r)
        r += 1
        count += 1

    if own:
        _save(wb)
    print(f"  {count} fórmulas VENDIDOS regeneradas sobre {len(months)} meses: {months}")


def add_dropdowns_to_sheet(ws, last_row=1000):
    """Aplica todos los dropdowns estándar a una hoja de ventas mensual.

    Estructura asumida (ABRIL 26 en adelante): C=PDV, D=PIEZA, F=METAL,
    I=FORMA PAGO, J=DESC STOCK, L=ESTADO.

    - PIEZA: tipos únicos tomados de la hoja STOCK (+ ASIGNAR)
    - METAL: ALPACA / BRONCE
    - PDV: LOCAL / ECOMMERCE / INSTAGRAM
    - FORMA PAGO: SOMOS.ARDE / GETNET / EFECTIVO / MERCADO PAGO / TRANSFERENCIA
    - DESC STOCK: SI / NO
    - ESTADO: ARCHIVADO / EMPAQUETADO / POR ENVIAR / POR EMPAQUETAR / ASIGNAR
    """
    piezas = _stock_piezas(ws.parent)
    pieza_list = "ASIGNAR," + ",".join(piezas) if piezas else "ASIGNAR"

    specs = {
        "C": '"ASIGNAR,LOCAL,ECOMMERCE,INSTAGRAM"',
        "D": f'"{pieza_list}"',
        "F": '"ASIGNAR,ALPACA,BRONCE"',
        "I": '"ASIGNAR,SOMOS.ARDE,GETNET,EFECTIVO,MERCADO PAGO,TRANSFERENCIA"',
        "J": '"SI,NO"',
        "L": '"ASIGNAR,ARCHIVADO,EMPAQUETADO,POR ENVIAR,POR EMPAQUETAR"',
    }
    for col, f1 in specs.items():
        dv = DataValidation(type="list", formula1=f1,
                            allow_blank=True, showDropDown=False)
        dv.sqref = f"{col}2:{col}{last_row}"
        ws.add_data_validation(dv)

# ── Función: add_month_sheet ───────────────────────────────────────────────────

def add_month_sheet(month_name: str, wb=None):
    """Crea una hoja de mes nueva con la estructura completa (copia de AGOSTO 26)."""
    own = wb is None
    if own:
        wb = _load()

    if month_name in wb.sheetnames:
        print(f"  La hoja '{month_name}' ya existe.")
        if own:
            return
        return

    ref = "AGOSTO 26" if "AGOSTO 26" in wb.sheetnames else "JULIO 26"
    ws_ref = wb[ref]
    ws_new = wb.create_sheet(month_name)

    # Copiar header (fila 1)
    for cell in ws_ref[1]:
        nc = ws_new.cell(1, cell.column)
        nc.value = cell.value
        if cell.has_style:
            nc.font = copy(cell.font)
            nc.fill = copy(cell.fill)
            nc.alignment = copy(cell.alignment)

    # Copiar anchos de columna
    for col_letter, dim in ws_ref.column_dimensions.items():
        ws_new.column_dimensions[col_letter].width = dim.width

    # Dropdown DESC STOCK
    ds_col = DESC_STOCK_COL.get(month_name, "J")
    _add_desc_stock_dv(ws_new, ds_col)

    # Registrar la columna en el mapa para fix_vendidos_formulas
    DESC_STOCK_COL[month_name] = ds_col

    # Posicionar después de la última hoja de mes
    last_month_idx = max(
        (wb.sheetnames.index(s) for s in wb.sheetnames
         if "26" in s and s != month_name and s != "PAUYSOL ABRIL"),
        default=0
    )
    current_idx = wb.sheetnames.index(month_name)
    wb.move_sheet(month_name, offset=(last_month_idx - current_idx + 1))

    if own:
        _save(wb)
    print(f"  Hoja '{month_name}' creada desde '{ref}'.")

# ── Función: stock_report ──────────────────────────────────────────────────────

def stock_report(wb=None):
    """Imprime un resumen del estado del stock calculando VENDIDOS desde las hojas de ventas."""
    own = wb is None
    if own:
        wb = _load()

    ws_stock = wb["STOCK"]
    hr = _find_header_row(ws_stock)

    # Preconstruir índice de ventas: {(pieza, modelo, metal): count}
    sold = {}
    for sheet_name, ds_col in DESC_STOCK_COL.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws_m = wb[sheet_name]
        ds_idx = column_index_from_string(ds_col)
        for row in ws_m.iter_rows(min_row=2, values_only=True):
            if not row[3]:  # col D = PIEZA
                continue
            if str(row[ds_idx - 1]).upper() == "SI":
                key = (row[3], row[4], row[5])  # PIEZA, MODELO, METAL
                sold[key] = sold.get(key, 0) + 1

    counts = {"OK": [], "STOCK BAJO": [], "SIN STOCK": [], "Sin cargar": []}

    r = hr + 1
    while ws_stock.cell(r, _col("pieza")).value:
        pieza    = ws_stock.cell(r, _col("pieza")).value
        modelo   = ws_stock.cell(r, _col("modelo")).value
        metal    = ws_stock.cell(r, _col("metal")).value
        inicial  = ws_stock.cell(r, _col("stock_inicial")).value
        perdidos = ws_stock.cell(r, _col("perdidos")).value or 0
        taller   = ws_stock.cell(r, _col("taller")).value or 0
        label = f"{pieza} {modelo} {metal}"

        if not inicial:
            counts["Sin cargar"].append(label)
        else:
            vendidos   = sold.get((pieza, modelo, metal), 0)
            stock_local = inicial - vendidos - perdidos - taller
            if stock_local <= 0:
                counts["SIN STOCK"].append(label)
            elif stock_local <= 3:
                counts["STOCK BAJO"].append(f"{label} → {stock_local} u.")
            else:
                counts["OK"].append(label)
        r += 1

    print("\n=== REPORTE DE STOCK ===")
    for k, v in counts.items():
        print(f"  {k}: {len(v)}")
    for section in ["SIN STOCK", "STOCK BAJO"]:
        if counts[section]:
            print(f"\n{section}:")
            for item in counts[section]:
                print(f"    {item}")

# ── Función: fix_desc_stock_values ─────────────────────────────────────────────

def fix_desc_stock_values(wb=None):
    """Convierte cualquier booleano True/False a 'SI'/'NO' en columnas DESC STOCK."""
    own = wb is None
    if own:
        wb = _load()

    changed = 0
    for sheet_name, col_letter in DESC_STOCK_COL.items():
        if sheet_name not in wb.sheetnames:
            continue
        col_idx = column_index_from_string(col_letter)
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            cell = row[0]
            if cell.value is True:
                cell.value = "SI"
                changed += 1
            elif cell.value is False:
                cell.value = "NO"
                changed += 1

    if own:
        _save(wb)
    print(f"  {changed} valores DESC STOCK convertidos a SI/NO.")

# ── Función: validate_entries ───────────────────────────────────────────────────

def _stock_combos(wb):
    """Devuelve (set_de_tuplas, lista_de_strings) de (PIEZA, MODELO, METAL) de STOCK."""
    st = wb["STOCK"]
    hr = _find_header_row(st)
    combos, strings = set(), []
    r = hr + 1
    while st.cell(r, _col("pieza")).value:
        p = st.cell(r, _col("pieza")).value
        m = st.cell(r, _col("modelo")).value
        me = st.cell(r, _col("metal")).value
        key = (
            str(p).strip().upper() if p is not None else "",
            str(m).strip().upper() if m is not None else "",
            str(me).strip().upper() if me is not None else "",
        )
        if key not in combos:
            combos.add(key)
            strings.append(" | ".join(key))
        r += 1
    return combos, strings

def validate_entries(fix=False, wb=None):
    """Detecta entradas en hojas mensuales donde PIEZA+MODELO+METAL no coincide con STOCK.

    Usa difflib para sugerir la corrección más probable.
    Si fix=True, aplica las correcciones automáticamente.
    """
    own = wb is None
    if own:
        wb = _load()

    combos, combo_strings = _stock_combos(wb)
    str_to_tuple = {s: tuple(s.split(" | ")) for s in combo_strings}

    problems = []   # (hoja, fila, actual, sugerencia)
    fixed = 0

    for name in _month_sheets(wb):
        ws = wb[name]
        ds_col = column_index_from_string(DESC_STOCK_COL.get(name, "J"))
        # PIEZA=D(4), MODELO=E(5), METAL=F(6) en todas las hojas mensuales
        r = 2
        last = ws.max_row
        while r <= last:
            pieza = ws.cell(r, 4).value
            if pieza is None:
                r += 1
                continue
            desc = ws.cell(r, ds_col).value
            if str(desc).strip().upper() != "SI":
                r += 1
                continue
            modelo = ws.cell(r, 5).value
            metal = ws.cell(r, 6).value
            key = (
                str(pieza).strip().upper(),
                str(modelo).strip().upper() if modelo is not None else "",
                str(metal).strip().upper() if metal is not None else "",
            )
            if key in combos:
                r += 1
                continue
            actual = " | ".join(key)
            match = difflib.get_close_matches(actual, combo_strings, n=1, cutoff=0.6)
            suggestion = match[0] if match else None
            problems.append((name, r, actual, suggestion))
            if fix and suggestion:
                sp, sm, sme = str_to_tuple[suggestion]
                ws.cell(r, 4).value = sp
                ws.cell(r, 5).value = sm
                ws.cell(r, 6).value = sme
                fixed += 1
            r += 1

    print(f"\n=== VALIDACIÓN DE ENTRADAS (DESC STOCK = SI) ===")
    print(f"  Combinaciones válidas en STOCK: {len(combos)}")
    print(f"  Entradas con posibles typos: {len(problems)}")
    for hoja, fila, actual, sug in problems:
        arrow = f"  ->  {sug}" if sug else "  ->  (sin sugerencia)"
        print(f"    [{hoja}] fila {fila}: {actual}{arrow}")
    if fix:
        if own:
            _save(wb)
        print(f"  {fixed} entradas corregidas automáticamente.")
    return problems

# ── CLI ────────────────────────────────────────────────────────────────────────

COMMANDS = {
    "styling":       (apply_styling,        "Aplicar colores a todas las hojas"),
    "fix_formulas":  (fix_vendidos_formulas, "Regenerar fórmulas VENDIDOS en STOCK"),
    "add_month":     (None,                  "Crear nueva hoja de mes (requiere nombre)"),
    "stock_report":  (stock_report,          "Ver estado del stock por consola"),
    "fix_desc_stock":(fix_desc_stock_values, "Convertir TRUE/FALSE a SI/NO en DESC STOCK"),
    "validate":      (None,                   "Detectar (y opcional --fix) entradas con typos vs STOCK"),
}

def main():
    args = sys.argv[1:]

    if not args:
        print(__doc__)
        print("Comandos disponibles:")
        for cmd, (_, desc) in COMMANDS.items():
            print(f"  {cmd:<18} {desc}")
        return

    cmd = args[0]

    if cmd == "add_month":
        if len(args) < 2:
            print("Uso: python arde_tools.py add_month 'SEPTIEMBRE 26'")
        else:
            add_month_sheet(args[1])
    elif cmd == "validate":
        validate_entries(fix=("--fix" in args))
    elif cmd in COMMANDS and COMMANDS[cmd][0]:
        COMMANDS[cmd][0]()
    else:
        print(f"Comando desconocido: '{cmd}'")
        print("Corré sin argumentos para ver la lista de comandos.")

if __name__ == "__main__":
    main()
